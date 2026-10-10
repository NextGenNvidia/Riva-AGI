"""Unit test suite for Universal Ingestion components."""

import json
from pathlib import Path
from unittest.mock import MagicMock
import pytest

from rag_knowledge.ingestion.privacy import (
    PrivacyGateError,
    assert_no_privacy_leaks,
    redact_private_text,
    scan_for_privacy_leaks,
)
from rag_knowledge.ingestion.joiner import StudentEntityJoiner
from rag_knowledge.ingestion.readers.word import WordDocxReader
from rag_knowledge.storage.qdrant_storage import QdrantKnowledgeStore


def test_unanchored_privacy_gate_detection():
    # Test phone embedded inside sentence
    text_with_phone = "Contact the student coordinator at +91 9876543210 for registration."
    leaks = scan_for_privacy_leaks(text_with_phone)
    assert len(leaks["phones"]) == 1
    assert "9876543210" in leaks["phones"][0]

    # Test email embedded inside sentence
    text_with_email = "Send notes to student.help@university.edu before noon."
    leaks = scan_for_privacy_leaks(text_with_email)
    assert len(leaks["emails"]) == 1
    assert leaks["emails"][0] == "student.help@university.edu"

    # Test fail-closed assertion
    doc = {
        "id": "doc_leak",
        "title": "Admissions",
        "content": "Call 9876543210 immediately",
    }
    with pytest.raises(PrivacyGateError):
        assert_no_privacy_leaks([doc])


def test_safe_identifier_whitelist():
    # Safe 10+ digit UID or Roll Number should NOT trigger a phone leak
    safe_roll = "2025R0111101103"
    text_with_roll = f"Student with Roll Number {safe_roll} has completed registration."
    leaks = scan_for_privacy_leaks(text_with_roll, safe_identifiers={safe_roll})
    assert len(leaks["phones"]) == 0

    doc = {
        "id": "doc_safe",
        "title": "Student Record",
        "content": f"Roll number is {safe_roll}.",
    }
    # Should not raise
    assert_no_privacy_leaks([doc], safe_identifiers={safe_roll})


def test_opt_in_redaction():
    text = "Contact 9876543210 or email test@example.com for info."
    redacted = redact_private_text(text)
    assert "9876543210" not in redacted
    assert "test@example.com" not in redacted
    assert "[REDACTED_EMAIL]" in redacted
    assert "[REDACTED_PHONE]" in redacted


def test_word_docx_reader(tmp_path):
    import docx

    doc_file = tmp_path / "sample_policy.docx"
    doc = docx.Document()
    doc.add_heading("Academic Regulations", level=1)
    doc.add_paragraph("Students must maintain 75% attendance.")
    doc.add_heading("Leave Policy", level=2)
    doc.add_paragraph("Medical leaves require certificate.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Category"
    table.rows[0].cells[1].text = "Minimum"
    table.rows[1].cells[0].text = "Theory"
    table.rows[1].cells[1].text = "75%"

    doc.save(str(doc_file))

    reader = WordDocxReader()
    assert reader.can_read(doc_file)
    units = reader.read(doc_file)

    assert len(units) >= 2
    titles = [u["title"] for u in units]
    assert any("Academic Regulations" in t or "Leave Policy" in t for t in titles)
    assert any(u["unit_type"] == "table" for u in units)


def test_qdrant_batch_delete():
    store = QdrantKnowledgeStore(url="https://fake.qdrant.io", api_key="fake")
    mock_client = MagicMock()
    store._client = mock_client
    store._is_connected = True

    deleted_count = store.delete_documents(["doc_1", "doc_2", "doc_3"])
    assert deleted_count == 3
    assert mock_client.delete.called



def test_known_private_value_leak_detection():
    # Detect father's name or personal phone passed in memory
    private_names = {"Sanjeev Dutt", "Bhanu Pratap Rai"}
    doc = {
        "id": "doc_test_father",
        "title": "Nominal Roll Record",
        "content": "Student record with guardian Sanjeev Dutt listed.",
    }
    with pytest.raises(PrivacyGateError):
        assert_no_privacy_leaks([doc], known_private_values=private_names)


def test_full_field_redaction():
    doc = {
        "id": "doc_leak_all",
        "title": "Meeting with test@example.com",
        "summary": "Notes for +91 9876543210 coordinator",
        "content": "Secret note: contact 9876543210 or email test@example.com",
        "aliases": ["test@example.com"],
        "keywords": ["9876543210"],
    }
    assert_no_privacy_leaks([doc], opt_in_redact=True)
    assert "[REDACTED_EMAIL]" in doc["title"]
    assert "[REDACTED_PHONE]" in doc["summary"]
    assert "[REDACTED_EMAIL]" in doc["aliases"][0]
    assert "[REDACTED_PHONE]" in doc["keywords"][0]


def test_tabular_metadata_sanitization_and_chunking(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "candidates.csv"
    # Create 50 candidates in the same domain to trigger member list chunking (>40)
    lines = ["Name,Email,Phone,Domain"]
    for i in range(1, 51):
        lines.append(f"Candidate {i},candidate{i}@test.com,98765432{i:02d},AI/ML")
    csv_file.write_text("\n".join(lines), encoding="utf-8")

    docs = load_generic_tabular_dataset(csv_file)
    assert len(docs) > 0

    # Verify no private emails or phones in any document metadata
    for d in docs:
        meta = d.get("metadata", {})
        assert "private" not in meta
        for k, v in meta.items():
            assert "email" not in k.lower()
            assert "phone" not in k.lower()
            assert "@test.com" not in str(v)

    # Verify chunked aggregation docs exist for AI/ML (Part 1 and Part 2)
    chunked = [d for d in docs if "part" in d["id"].lower() and "ai_ml" in d["id"].lower()]
    assert len(chunked) == 2
    assert "Part 1/2" in chunked[0]["title"]
    assert "Part 2/2" in chunked[1]["title"]


def test_image_document_ingestion(tmp_path, monkeypatch):
    from rag_knowledge.ingestion.ingest import load_image_documents, load_source_documents

    img_file = tmp_path / "diagram.png"
    img_file.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")

    mock_text = (
        "# System Architecture\n"
        "The system coordinates query ingestion and vector retrieval.\n\n"
        "---\n\n"
        "## Performance Metrics\n"
        "Latency is under 20ms for vector lookup."
    )

    monkeypatch.setattr("rag_knowledge.ingestion.ingest.extract_image_text", lambda p: mock_text)

    docs = load_image_documents(img_file)
    assert len(docs) >= 2
    assert any(d["id"] == "img_diagram_full" for d in docs)
    assert any("Architecture" in d["title"] or "Performance" in d["title"] for d in docs)

    # Verify load_source_documents detects image extension
    auto_docs = load_source_documents(img_file)
    assert len(auto_docs) == len(docs)


def test_phone_with_space_in_waitlist_filtered(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "waitlist.csv"
    csv_file.write_text(
        "Name,Status,Phone,Father Name,Score\n"
        "Aarav Sharma,Waiting List,98765 43210,Rajesh Sharma,92\n"
        "Neha Verma,Waiting List,+91 91234-56789,Sanjay Verma,88\n",
        encoding="utf-8",
    )
    docs = load_generic_tabular_dataset(csv_file)
    wl_docs = [d for d in docs if "waiting_list" in d["id"]]
    assert len(wl_docs) == 1
    content = wl_docs[0]["content"]

    # Ensure no phone numbers (spaced or formatted) or father's name leaked into waitlist embedded text
    assert "98765 43210" not in content
    assert "91234" not in content
    assert "Rajesh Sharma" not in content
    assert "Sanjay Verma" not in content
    assert "Score: 92" in content


def test_long_roll_number_exempt_from_phone_drop(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "students.csv"
    # Roll number starting with 9 and >= 10 digits that might resemble a phone number
    long_roll = "98202510001"
    csv_file.write_text(
        "Candidate Name,Roll Number,Branch\n"
        f"Deepak Kumar,{long_roll},CSE\n",
        encoding="utf-8",
    )
    docs = load_generic_tabular_dataset(csv_file)
    rec_docs = [d for d in docs if d["category"] == "record"]
    assert len(rec_docs) == 1
    # Verify roll number was NOT dropped by the phone filter
    assert long_roll in rec_docs[0]["content"]
    assert rec_docs[0]["metadata"].get("Roll Number") == long_roll


def test_image_no_cloud_vision_skips_without_junk_record(tmp_path, monkeypatch):
    from rag_knowledge.ingestion.ingest import load_image_documents

    img_file = tmp_path / "scan.png"
    img_file.write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")

    # Simulate OCR failure when cloud vision is not allowed
    monkeypatch.setattr(
        "rag_knowledge.ingestion.ingest.extract_image_text",
        lambda p, allow_cloud_vision=False: ("", "none", 0.0),
    )

    docs = load_image_documents(img_file, allow_cloud_vision=False)
    # Must skip completely and NOT produce junk fallback "Image document scan.png"
    assert docs == []


def test_father_name_in_content_flagged_and_redacted():
    doc = {
        "id": "doc_eval",
        "title": "Nominal Roll Summary",
        "content": "Candidate: Rohan Gupta, Father Name: Ramesh Gupta, Branch: IT",
    }
    # 1. Flagged and blocked when unredacted
    with pytest.raises(PrivacyGateError):
        assert_no_privacy_leaks([dict(doc)], known_private_values={"Ramesh Gupta"})

    # 2. Successfully redacted when opt_in_redact is True
    cloned = dict(doc)
    assert_no_privacy_leaks([cloned], known_private_values={"Ramesh Gupta"}, opt_in_redact=True)
    assert "Ramesh Gupta" not in cloned["content"]
    assert "[REDACTED]" in cloned["content"]


def test_residence_phone_not_exempt_from_filter(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "candidates_residence.csv"
    # Columns 'Residence Phone' and 'Paid Contact' contain 'id' as a substring.
    # Neither must be treated as an ID column, and phone numbers must be strictly filtered.
    csv_file.write_text(
        "Candidate Name,Residence Phone,Paid Contact,Branch\n"
        "Shreya Singh,9876543210,9123456789,ECE\n",
        encoding="utf-8",
    )
    docs = load_generic_tabular_dataset(csv_file)
    rec_docs = [d for d in docs if d["category"] == "record"]
    assert len(rec_docs) == 1
    doc_content = rec_docs[0]["content"]
    doc_meta = rec_docs[0]["metadata"]

    # Verify neither 'Residence Phone' nor 'Paid Contact' phone numbers leaked
    assert "9876543210" not in doc_content
    assert "9123456789" not in doc_content
    assert "Residence Phone" not in doc_meta
    assert "Paid Contact" not in doc_meta


def test_phone_number_in_id_column_filtered_when_matching_contact(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "club_form.csv"
    # An applicant mistakenly typed their phone number into the 'Student ID' column
    # as well as the 'Phone Number' column.
    csv_file.write_text(
        "Candidate Name,Student ID,Phone Number,Branch\n"
        "Test Applicant,9876543210,9876543210,ELCE\n",
        encoding="utf-8",
    )
    docs = load_generic_tabular_dataset(csv_file)
    rec_docs = [d for d in docs if d["category"] == "record"]
    assert len(rec_docs) == 1
    doc_content = rec_docs[0]["content"]
    doc_meta = rec_docs[0]["metadata"]

    # Verify the phone number was dropped from the ID column and not stored in metadata
    assert "9876543210" not in doc_content
    assert "Student ID" not in doc_meta
    assert "Phone Number" not in doc_meta


def test_word_document_ingestion(tmp_path):
    import docx
    from rag_knowledge.ingestion.ingest import load_source_documents

    doc_file = tmp_path / "club_handbook.docx"
    doc = docx.Document()
    doc.add_heading("Society Guidelines", level=1)
    doc.add_paragraph("Welcome to the AI & Robotics Society handbook for members.")

    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Role"
    hdr_cells[1].text = "Lead"
    hdr_cells[2].text = "Domain"

    row_cells = table.rows[1].cells
    row_cells[0].text = "President"
    row_cells[1].text = "Ankit Singh"
    row_cells[2].text = "Core"

    doc.save(str(doc_file))

    docs = load_source_documents(doc_file)
    assert len(docs) >= 2
    prose_docs = [d for d in docs if "Society Guidelines" in d["title"]]
    assert len(prose_docs) >= 1
    assert "AI & Robotics Society" in prose_docs[0]["content"]

    table_docs = [d for d in docs if "Table" in d["title"]]
    assert len(table_docs) >= 1
    assert "President" in table_docs[0]["content"]
    assert "Core" in table_docs[0]["content"]


def test_underscore_header_roll_number_not_dropped(tmp_path):
    from rag_knowledge.ingestion.ingest import load_generic_tabular_dataset

    csv_file = tmp_path / "students_underscore.csv"
    long_roll = "210097010045"
    csv_file.write_text(
        "candidate_name,roll_number,student_branch\n"
        f"Deepak Kumar,{long_roll},CSE\n",
        encoding="utf-8",
    )
    docs = load_generic_tabular_dataset(csv_file)
    rec_docs = [d for d in docs if d["category"] == "record"]
    assert len(rec_docs) == 1
    assert long_roll in rec_docs[0]["content"]
    assert rec_docs[0]["metadata"].get("roll_number") == long_roll



