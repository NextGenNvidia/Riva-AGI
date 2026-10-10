"""Unit tests for PDF and Word loaders in PR 1b."""

from unittest.mock import MagicMock, patch
import pytest

from rag_knowledge.ingestion.loaders.dispatch import load_source_documents
from rag_knowledge.ingestion.loaders.pdf import load_pdf_documents
from rag_knowledge.ingestion.readers.word import WordDocxReader


def test_word_docx_reader(tmp_path):
    import docx

    doc_file = tmp_path / "sample_policy.docx"
    doc = docx.Document()
    doc.add_heading("Academic Regulations", level=1)
    doc.add_paragraph("Students must maintain 75% attendance.")
    doc.add_heading("Leave Policy", level=2)
    doc.add_paragraph("Medical leaves require certificate.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Category"
    table.rows[0].cells[1].text = "Minimum"
    table.rows[1].cells[0].text = "Theory"
    table.rows[1].cells[1].text = "75%"

    doc.save(str(doc_file))

    reader = WordDocxReader()
    assert reader.can_read(doc_file)
    units = reader.read(doc_file)

    assert len(units) >= 2
    titles = [u["title"] for u in units]
    assert any("Academic Regulations" in t or "Leave Policy" in t for t in titles)
    assert any(u["unit_type"] == "table" for u in units)


def test_word_document_ingestion(tmp_path):
    import docx

    doc_file = tmp_path / "club_handbook.docx"
    doc = docx.Document()
    doc.add_heading("Society Guidelines", level=1)
    doc.add_paragraph("Welcome to the AI & Robotics Society handbook for members.")

    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Role"
    hdr_cells[1].text = "Lead"
    hdr_cells[2].text = "Domain"

    row_cells = table.rows[1].cells
    row_cells[0].text = "President"
    row_cells[1].text = "Ankit Singh"
    row_cells[2].text = "Core"

    doc.save(str(doc_file))

    docs = load_source_documents(doc_file)
    assert len(docs) >= 2
    prose_docs = [d for d in docs if "Society Guidelines" in d["title"]]
    assert len(prose_docs) >= 1
    assert "AI & Robotics Society" in prose_docs[0]["content"]

    table_docs = [d for d in docs if "Table" in d["title"]]
    assert len(table_docs) >= 1
    assert "President" in table_docs[0]["content"]
    assert "Core" in table_docs[0]["content"]


def test_pdf_document_ingestion(tmp_path):
    with patch("pypdf.PdfReader") as mock_pdf:
        mock_page = MagicMock()
        mock_page.extract_text.return_value = "Course Overview\nIntroduction to Artificial Intelligence"
        mock_pdf.return_value.pages = [mock_page]

        pdf_file = tmp_path / "syllabus.pdf"
        pdf_file.write_bytes(b"%PDF-1.4\nfake pdf content")

        docs = load_pdf_documents(pdf_file)
        assert len(docs) == 1
        assert "Course Overview" in docs[0]["title"] or "syllabus" in docs[0]["title"].lower()
        assert "Artificial Intelligence" in docs[0]["content"]
        assert docs[0]["metadata"].get("source") == "syllabus.pdf"
        assert docs[0]["metadata"].get("page") == 1

        # Test dispatch loader handles pdf
        auto_docs = load_source_documents(pdf_file)
        assert len(auto_docs) == 1
