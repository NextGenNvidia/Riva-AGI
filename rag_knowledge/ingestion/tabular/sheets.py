"""Spreadsheet and delimited tabular file loaders."""

import csv
import logging
from pathlib import Path
from typing import List, Tuple

import openpyxl

logger = logging.getLogger("rag.ingest")


def open_tabular_sheets(filepath: Path) -> List[Tuple[str, List[str], List[List[str]]]]:
    """Opens CSV, TSV, or XLSX tabular sheets and extracts raw header and data rows.

    Rejects legacy binary .xls files with an informative warning.
    Returns a list of (sheet_title, raw_headers, data_rows) tuples.
    """
    if not filepath.is_file():
        return []

    ext = filepath.suffix.lower()
    table_name = filepath.stem.replace("_", " ").replace("-", " ").title()
    sheet_datasets: List[Tuple[str, List[str], List[List[str]]]] = []

    if ext in (".csv", ".tsv"):
        delimiter = "\t" if ext == ".tsv" else ","
        rows: List[List[str]] = []
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f, delimiter=delimiter)
            for r in reader:
                if any(r):
                    rows.append([str(c or "").strip() for c in r])
        if len(rows) >= 2:
            sheet_datasets.append((table_name, rows[0], rows[1:]))

    elif ext in (".xlsx", ".xls"):
        if ext == ".xls":
            logger.warning(
                f"File {filepath.name} is a legacy binary Excel format (.xls). "
                "openpyxl does not support binary .xls. Please convert to .xlsx or .csv."
            )
            return []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sname in wb.sheetnames:
                sheet = wb[sname]
                s_rows: List[List[str]] = []
                for r in sheet.iter_rows(values_only=True):
                    if any(c is not None and str(c).strip() for c in r):
                        s_rows.append([str(c or "").strip() if c is not None else "" for c in r])
                if len(s_rows) >= 2:
                    sub_title = table_name if len(wb.sheetnames) == 1 else f"{table_name} ({sname})"
                    sheet_datasets.append((sub_title, s_rows[0], s_rows[1:]))
            wb.close()
        except Exception as e:
            logger.warning(f"Could not read workbook {filepath.name}: {e}")
            return []

    return sheet_datasets
